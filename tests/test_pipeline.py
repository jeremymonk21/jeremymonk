from __future__ import annotations

import csv
from datetime import timedelta

import httpx
import pytest
import respx
from openpyxl import load_workbook

from monkleads import pipeline as pl
from monkleads.audit.pagespeed import PSI_URL
from monkleads.compliance import SuppressionList
from monkleads.config import AppConfig, Secrets
from monkleads.db import Audit, Business, OutreachStatus, find_business, set_status, utcnow
from monkleads.export import safe_cell, write_csv, write_xlsx
from monkleads.pipeline import Pipeline
from monkleads.report import render_report

GOOD_HTML = """<html><head><title>Good Co</title><meta name="description" content="x">
<meta name="viewport" content="width=device-width"></head><body>
<a href="tel:5185550000">Call</a><form><textarea></textarea></form>
<iframe src="https://www.google.com/maps/embed"></iframe><p>Mon-Fri 9am - 5pm</p>
<footer>© 2026</footer></body></html>"""


async def fake_resolver(host: str) -> list[object]:
    return [("ok",)]


@pytest.fixture
def pipe(cfg: AppConfig, secrets: Secrets, monkeypatch: pytest.MonkeyPatch) -> Pipeline:
    # Use a resolver that never touches real DNS.
    orig_init = pl.SiteChecker.__init__

    def init(self, client, acfg, **kw):  # type: ignore[no-untyped-def]
        kw["resolver"] = fake_resolver
        orig_init(self, client, acfg, **kw)

    monkeypatch.setattr(pl.SiteChecker, "__init__", init)
    p = Pipeline(cfg, secrets)
    with p.sessions() as s, s.begin():
        s.add_all(
            [
                Business(
                    place_id="no_site",
                    name='=HYPERLINK("evil")',
                    category_key="barber",
                    rating=4.7,
                    review_count=60,
                    distance_miles=1.0,
                    town="Troy",
                ),
                Business(
                    place_id="good",
                    name="Good Co",
                    category_key="restaurant",
                    website="https://good.example",
                    rating=4.2,
                    review_count=30,
                    distance_miles=2.0,
                    town="Troy",
                ),
                Business(
                    place_id="social",
                    name="Sal's",
                    category_key="salon",
                    website="https://facebook.com/sals",
                    rating=4.9,
                    review_count=80,
                    distance_miles=3.0,
                ),
                Business(
                    place_id="chain", name="Dunkin'", excluded_reason="chain: blocklisted chain (dunkin)"
                ),
                Business(place_id="won", name="Already A Client", category_key="cafe"),
                Business(place_id="dnc", name="Leave Me Alone LLC", category_key="cafe"),
            ]
        )
    with p.sessions() as s, s.begin():
        set_status(s, find_business(s, "won"), OutreachStatus.WON, "signed Standard")  # type: ignore[arg-type]
        set_status(s, find_business(s, "dnc"), OutreachStatus.DO_NOT_CONTACT)  # type: ignore[arg-type]
    return p


def mock_good_site() -> None:
    respx.get("https://good.example/robots.txt").respond(404)
    respx.get("http://good.example/robots.txt").respond(404)
    respx.get("https://good.example/").respond(200, html=GOOD_HTML)
    respx.get("http://good.example/").respond(301, headers={"location": "https://good.example/"})
    respx.get(PSI_URL).respond(
        200, json={"lighthouseResult": {"categories": {"performance": {"score": 0.35}}}}
    )


@respx.mock
async def test_full_flow_audit_score_export(pipe: Pipeline, cfg: AppConfig, tmp_path) -> None:  # type: ignore[no-untyped-def]
    mock_good_site()
    with pipe.sessions() as s:
        todo = pipe.select_for_audit(s)
    assert {b.place_id for b in todo} == {"no_site", "good", "social"}  # chain, won, dnc skipped

    stats = await pipe.audit(todo, run_psi=True, screenshots=False)
    assert stats.audited == 3 and stats.errors == 0 and stats.pagespeed == 1
    assert stats.by_status == {"none": 1, "ok": 1, "social_only": 1}

    sstats = await pipe.score(use_llm=False)
    assert sstats.scored == 3 and sstats.email_drafts == 3

    rows, suppressed = pipe.ranked_leads()
    assert suppressed == 0
    assert [r.place_id for r in rows][0] == "no_site"  # no website ranks first
    good = next(r for r in rows if r.place_id == "good")
    assert good.performance == 35 and any("phone" in p for p in good.problems)
    assert all(r.email_body and "123 Example St" in r.email_body for r in rows)

    # Suppression is applied before export
    SuppressionList(cfg.output.suppression_file).add("domain", "good.example")
    rows2, suppressed2 = pipe.ranked_leads()
    assert suppressed2 == 1 and "good" not in {r.place_id for r in rows2}

    csv_path = write_csv(rows2, tmp_path / "out.csv")
    with csv_path.open(encoding="utf-8-sig") as fh:
        data = list(csv.DictReader(fh))
    assert data[0]["Business"].startswith("'=")  # formula injection neutralised
    assert int(data[0]["Score"]) >= int(data[-1]["Score"])

    xlsx_path = write_xlsx(rows2, tmp_path / "out.xlsx")
    wb = load_workbook(xlsx_path)
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    maps_col = headers.index("Google Maps") + 1
    assert ws.cell(row=2, column=headers.index("Score") + 1).value == rows2[0].score
    assert ws.conditional_formatting  # score colour scale etc.
    assert ws.freeze_panes == "E2"
    assert maps_col  # hyperlink column exists

    html = render_report(
        rows2, tmp_path / "report.html", sender=cfg.sender, suppressed=suppressed2
    ).read_text()
    assert "Top 2 leads" in html and "1 suppressed" in html
    assert "&#34;evil&#34;" in html or "&quot;evil&quot;" in html  # autoescaped


@respx.mock
async def test_incremental_audit_only_stale(pipe: Pipeline, cfg: AppConfig) -> None:
    mock_good_site()
    with pipe.sessions() as s:
        todo = pipe.select_for_audit(s)
    await pipe.audit(todo, run_psi=False, screenshots=False)
    with pipe.sessions() as s:
        assert pipe.select_for_audit(s) == []
        assert len(pipe.select_for_audit(s, force=True)) == 3

    # Age one audit past the stale window
    with pipe.sessions() as s, s.begin():
        b = find_business(s, "good")
        assert b is not None
        b.latest_audit.created_at = utcnow() - timedelta(days=cfg.audit.stale_after_days + 1)  # type: ignore[union-attr]
    with pipe.sessions() as s:
        assert [b.place_id for b in pipe.select_for_audit(s)] == ["good"]

    # A changed website triggers a re-audit even if the audit is fresh
    with pipe.sessions() as s, s.begin():
        b = find_business(s, "social")
        assert b is not None
        b.website = "https://sals-salon.example"
    with pipe.sessions() as s:
        assert "social" in {b.place_id for b in pipe.select_for_audit(s)}


@respx.mock
async def test_status_changes_hide_leads(pipe: Pipeline) -> None:
    mock_good_site()
    with pipe.sessions() as s:
        todo = pipe.select_for_audit(s)
    await pipe.audit(todo, run_psi=False, screenshots=False)
    await pipe.score(use_llm=False)
    assert len(pipe.ranked_leads()[0]) == 3

    with pipe.sessions() as s, s.begin():
        set_status(s, find_business(s, "good"), OutreachStatus.CONTACTED, "left a voicemail")  # type: ignore[arg-type]
    rows, _ = pipe.ranked_leads()
    good = next(r for r in rows if r.place_id == "good")
    assert good.status == "contacted" and "left a voicemail" in good.notes

    with pipe.sessions() as s, s.begin():
        set_status(s, find_business(s, "good"), OutreachStatus.WON)  # type: ignore[arg-type]
    assert "good" not in {r.place_id for r in pipe.ranked_leads()[0]}
    with pipe.sessions() as s:
        assert "good" not in {b.place_id for b in pipe.select_for_audit(s, force=True)}


async def test_score_without_sender_address_skips_email(pipe: Pipeline, cfg: AppConfig) -> None:
    cfg.sender.physical_address = ""
    with pipe.sessions() as s, s.begin():
        b = find_business(s, "no_site")
        assert b is not None
        s.add(Audit(business=b, site_status="none", findings={"input_url": None, "site_status": "none"}))
    stats = await pipe.score(use_llm=False)
    assert stats.scored == 1 and stats.email_drafts == 0
    assert any("physical_address" in w for w in stats.warnings)
    rows, _ = pipe.ranked_leads()
    assert rows[0].opener and rows[0].email_body is None


def test_safe_cell() -> None:
    assert safe_cell("=1+1") == "'=1+1"
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)"
    assert safe_cell("Joe's") == "Joe's"
    assert safe_cell(42) == 42


def test_cost_estimates(cfg: AppConfig) -> None:
    towns, cats = cfg.select_towns(["Troy", "Nassau"]), cfg.select_categories(["cafe", "barber", "plumber"])
    est = pl.estimate_discovery_cost(cfg, towns, cats)
    line = est.lines[0]
    assert line.calls_low == 6 and line.calls_high == 6 * cfg.discovery.max_pages_per_query
    assert est.high == pytest.approx(line.calls_high * cfg.costs.places_text_search_per_1000 / 1000)
    llm = pl.estimate_opener_cost(cfg, 10)
    assert llm.low == 0 and llm.high == pytest.approx(10 * pl.estimate_llm_cost(1, cfg.llm))


def test_http_client_sends_user_agent(pipe: Pipeline) -> None:
    client = pipe._http(5)
    assert isinstance(client, httpx.AsyncClient)
    assert client.headers["User-Agent"].startswith("MonkPressSiteCheck")
