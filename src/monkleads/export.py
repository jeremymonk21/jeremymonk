"""CSV and formatted Excel exports of the ranked lead list."""

from __future__ import annotations

import csv
import os
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .pipeline import LeadRow

COLUMNS: list[tuple[str, int]] = [
    ("Rank", 6),
    ("ID", 6),
    ("Score", 7),
    ("Business", 30),
    ("Category", 14),
    ("Town", 16),
    ("Phone", 15),
    ("Rating", 7),
    ("Reviews", 8),
    ("Miles", 7),
    ("Site status", 14),
    ("Mobile speed", 9),
    ("Platform", 14),
    ("Suggested plan", 14),
    ("Why this plan", 40),
    ("Problem 1", 50),
    ("Problem 2", 50),
    ("Problem 3", 50),
    ("Opener", 70),
    ("Status", 12),
    ("Notes", 30),
    ("Website", 30),
    ("Google Maps", 14),
    ("Screenshot", 14),
    ("Score breakdown", 60),
    ("Email subject", 30),
    ("Email draft", 70),
    ("Address", 40),
]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection in text that came from the internet."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _site_label(row: LeadRow) -> str:
    if row.site_status == "broken" and row.broken_reason:
        return f"broken ({row.broken_reason.replace('_', ' ')})"
    return {"none": "no website", "social_only": "social only"}.get(row.site_status, row.site_status)


def _breakdown(row: LeadRow) -> str:
    return "; ".join(f"{r['points']:+g} {r['detail']}" for r in row.reasons)


def _values(rank: int, row: LeadRow) -> list[Any]:
    probs = (row.problems + ["", "", ""])[:3]
    return [
        rank,
        row.id,
        row.score,
        row.name,
        row.category,
        row.town,
        row.phone,
        row.rating,
        row.review_count,
        row.distance_miles,
        _site_label(row),
        row.performance,
        row.platform,
        row.tier,
        row.tier_reason,
        *probs,
        row.opener,
        row.status,
        row.notes,
        row.website,
        row.maps_url,
        str(row.screenshot) if row.screenshot else None,
        _breakdown(row),
        row.email_subject,
        row.email_body,
        row.address,
    ]


def write_csv(rows: Sequence[LeadRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel reads UTF-8
        writer = csv.writer(fh)
        writer.writerow([c for c, _ in COLUMNS])
        for i, row in enumerate(rows, 1):
            writer.writerow([safe_cell(v) for v in _values(i, row)])
    return path


def _link(ws: Worksheet, cell_ref: str, target: str | None, label: str) -> None:
    cell = ws[cell_ref]
    if not target:
        cell.value = None
        return
    cell.value = label
    cell.hyperlink = target
    cell.font = Font(color="1F5FBF", underline="single")


def write_xlsx(rows: Sequence[LeadRow], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Leads"
    headers = [c for c, _ in COLUMNS]
    col = {name: get_column_letter(i + 1) for i, name in enumerate(headers)}
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for i, (_, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    wrap_cols = {
        "Why this plan",
        "Problem 1",
        "Problem 2",
        "Problem 3",
        "Opener",
        "Notes",
        "Score breakdown",
        "Email draft",
        "Address",
        "Business",
    }

    for i, row in enumerate(rows, 1):
        r = i + 1
        ws.append([safe_cell(v) for v in _values(i, row)])
        for name, letter in col.items():
            ws[f"{letter}{r}"].alignment = wrap if name in wrap_cols else top
        _link(ws, f"{col['Website']}{r}", row.website, row.website or "")
        _link(ws, f"{col['Google Maps']}{r}", row.maps_url, "Open map")
        if row.screenshot:
            rel = os.path.relpath(row.screenshot.resolve(), path.parent.resolve())
            _link(ws, f"{col['Screenshot']}{r}", rel, "View")
        ws.row_dimensions[r].height = 90

    last = len(rows) + 1
    if rows:
        score_rng = f"{col['Score']}2:{col['Score']}{last}"
        ws.conditional_formatting.add(
            score_rng,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F8696B",
                mid_type="num",
                mid_value=50,
                mid_color="FFEB84",
                end_type="num",
                end_value=100,
                end_color="63BE7B",
            ),
        )
        speed_rng = f"{col['Mobile speed']}2:{col['Mobile speed']}{last}"
        ws.conditional_formatting.add(
            speed_rng,
            CellIsRule(operator="lessThan", formula=["50"], fill=PatternFill("solid", fgColor="F8CBAD")),
        )
        site_col = col["Site status"]
        site_rng = f"{site_col}2:{site_col}{last}"
        for prefix, color in (("no website", "F8CBAD"), ("broken", "F4B084"), ("social only", "FFE699")):
            ws.conditional_formatting.add(
                site_rng,
                FormulaRule(
                    formula=[f'LEFT(${site_col}2,{len(prefix)})="{prefix}"'],
                    fill=PatternFill("solid", fgColor=color),
                ),
            )

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(last, 1)}"

    # Second sheet: how scoring works, for transparency.
    info = wb.create_sheet("How to read")
    info.column_dimensions["A"].width = 110
    for line in (
        "Score: 0-100, higher = bigger opportunity. 'Score breakdown' lists every point added or removed.",
        "Site status: no website / social only / broken (dns, parked, error, coming soon...) / ok / blocked (site refused automated checks).",
        "Mobile speed: Google PageSpeed Insights mobile performance score (0-100). Under 50 is poor.",
        "Problems are written in plain English for calls, visits, and emails. Openers are drafts: read and edit before using.",
        "Nothing is ever sent automatically. Leads marked won or do_not_contact, and anything on the suppression list, are excluded.",
        "Business listing data (name, address, phone, rating, reviews, hours) comes from Google (Places API).",
    ):
        info.append([line])
    wb.save(path)
    return path
